import openpyxl
import json
import urllib.parse
import os

wb = openpyxl.load_workbook("d:/IIM A/Study Material/Resources for lectures.xlsx")
sheet = wb['Sheet1']

# Helper to clean Google redirect URLs
def clean_url(url):
    if not url:
        return None
    url = str(url).strip()
    if "google.com/url?q=" in url:
        parsed = urllib.parse.urlparse(url)
        queries = urllib.parse.parse_qs(parsed.query)
        if 'q' in queries:
            return queries['q'][0]
    return url

sections = []
current_section = None
current_category = None

for r in range(1, sheet.max_row + 1):
    id_val = sheet.cell(row=r, column=1).value
    name_val = sheet.cell(row=r, column=2).value
    
    if not name_val:
        continue
        
    name_val = str(name_val).strip()
    
    # Detect Section Headers
    is_header = False
    if id_val is None or (isinstance(id_val, str) and not id_val.strip().replace('.', '').isdigit()):
        is_header = True
    elif "-" in name_val or "Gejo Speaks" in name_val or "Supergrads" in name_val or "Ravi Prakash" in name_val:
        is_header = True
        
    if name_val.startswith("Quants -") or name_val.startswith("VARC -") or name_val.startswith("DILR -") or name_val == "VARC - Gejo Speaks" or name_val == "VARC - Supergrads" or name_val == "VARC - Grok RCs" or name_val == "VARC - Ravi Prakash":
        is_header = True
        
    if is_header:
        # Ignore dummy letters or very short text headers
        if len(name_val) <= 2:
            continue
            
        current_section = name_val
        current_category = "Quants" if "quants" in name_val.lower() else "VARC" if "varc" in name_val.lower() else "DILR" if "dilr" in name_val.lower() else "General"
        sections.append({
            "section": current_section,
            "category": current_category,
            "topics": []
        })
        continue

    # If it is a topic row
    topic_name = name_val
    lectures = []
    
    for c in range(3, sheet.max_column + 1):
        cell = sheet.cell(row=r, column=c)
        if cell.value:
            lec_name = str(cell.value).strip()
            raw_url = cell.hyperlink.target if cell.hyperlink else None
            if not raw_url and "http" in lec_name:
                raw_url = lec_name
                lec_name = f"Lecture {c-2}"
            
            cleaned_url = clean_url(raw_url)
            if cleaned_url:
                lectures.append({
                    "name": lec_name,
                    "url": cleaned_url
                })
                
    if lectures:
        if not current_section:
            current_section = "General Prep Lectures"
            current_category = "General"
            sections.append({
                "section": current_section,
                "category": current_category,
                "topics": []
            })
            
        sections[-1]["topics"].append({
            "topic": topic_name,
            "lectures": lectures
        })

# Filter out sections with 0 topics
filtered_sections = [sec for sec in sections if len(sec["topics"]) > 0]

# Write to JSON
os.makedirs("d:/IIM A/src/data", exist_ok=True)
with open("d:/IIM A/src/data/video_lectures.json", "w", encoding="utf-8") as f:
    json.dump(filtered_sections, f, indent=2, ensure_ascii=False)

print("\nSuccessfully parsed and generated cleaned src/data/video_lectures.json!")
print(f"Stats: {len(filtered_sections)} sections parsed.")
for sec in filtered_sections:
    print(f" - {sec['section']}: {len(sec['topics'])} topics")
